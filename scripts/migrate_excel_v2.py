"""Migrate a v1 guild_scores.xlsx to the v2 schema in-place.

v1 layout (legacy):
  correct_nickname | latest_ocr_nickname | confidence | known_aliases |
  review_reason | peak_gear_score | 裝評 YYYY-MM-DD HH:MM |
  貢獻 YYYY-MM-DD HH:MM | 活躍 YYYY-MM-DD HH:MM | ...

v2 layout:
  correct_nickname | latest_ocr_nickname | confidence | review_reason |
  裝備評分(最高) | YYYY-MM-DD | YYYY-MM-DD | ...

Migration rules:
  * For each player, group all "裝評 YYYY-MM-DD HH:MM" columns by date and
    keep the max — matches the new "one slot per day, take max" policy.
  * Drop known_aliases / 貢獻 / 活躍 columns entirely.
  * Preserve correct_nickname, latest_ocr_nickname, confidence, review_reason.
  * Take a backup before writing.

Usage::

    .venv\\Scripts\\python.exe scripts\\migrate_excel_v2.py
    .venv\\Scripts\\python.exe scripts\\migrate_excel_v2.py --path D:\\my\\guild_scores.xlsx
    .venv\\Scripts\\python.exe scripts\\migrate_excel_v2.py --dry-run
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
except Exception:
    pass

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from ro_gearsync.storage import GuildScoresWorkbook, PlayerRecord  # noqa: E402
from ro_gearsync.utils.paths import user_data_dir  # noqa: E402


# Matches legacy v1 capture column headers: "裝評 YYYY-MM-DD HH:MM" or
# "裝評 YYYY-MM-DD" (rare). Only gear-score columns are migrated.
_LEGACY_GEAR_RE = re.compile(
    r"^裝評\s+(?P<date>\d{4}-\d{2}-\d{2})(?:\s+\d{2}:\d{2})?$"
)


def _is_legacy_header(headers: list) -> bool:
    """True if the sheet looks like the v1 layout."""
    text = " ".join(str(h) for h in headers if h)
    return "known_aliases" in text or "裝評 " in text or "peak_gear_score" in text


def migrate(path: Path, *, dry_run: bool = False) -> int:
    if not path.is_file():
        print(f"!! {path} not found")
        return 2
    print(f"[1] reading legacy workbook: {path}")
    book = load_workbook(path, data_only=True)
    if "Members" not in book.sheetnames:
        print(f"!! sheet 'Members' missing in {path}")
        return 3
    ws = book["Members"]
    headers = [c.value for c in ws[1]]
    if not _is_legacy_header(headers):
        print(
            "[*] this workbook doesn't look like the v1 schema — "
            "no migration needed."
        )
        return 0

    col_idx: dict[str, int] = {}
    for i, h in enumerate(headers):
        if h is None:
            continue
        col_idx[str(h).strip()] = i

    gear_columns: list[tuple[int, str]] = []  # (col_idx, date)
    for h, i in col_idx.items():
        m = _LEGACY_GEAR_RE.match(h)
        if m:
            gear_columns.append((i, m.group("date")))
    days = sorted({d for _, d in gear_columns})
    print(f"    legacy gear columns: {len(gear_columns)}  unique days: {len(days)}")

    new_wb = GuildScoresWorkbook(path)
    new_wb.capture_days = days

    n_rows = 0
    n_with_truth = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue

        def get(name: str) -> object | None:
            idx = col_idx.get(name)
            return row[idx] if idx is not None and idx < len(row) else None

        confidence_raw = get("confidence")
        try:
            confidence = (
                float(confidence_raw) if confidence_raw not in (None, "") else None
            )
        except (TypeError, ValueError):
            confidence = None

        rec = PlayerRecord(
            correct_nickname=str(get("correct_nickname") or "").strip(),
            latest_ocr_nickname=str(get("latest_ocr_nickname") or "").strip(),
            confidence=confidence,
            review_reason=str(get("review_reason") or "").strip(),
        )
        # Collapse legacy "裝評 YYYY-MM-DD HH:MM" columns to one value
        # per day (max), per the new schema.
        per_day_max: dict[str, int] = {}
        for idx, day in gear_columns:
            if idx >= len(row):
                continue
            cell = row[idx]
            if cell is None or cell == "":
                continue
            try:
                val = int(cell)
            except (TypeError, ValueError):
                continue
            prior = per_day_max.get(day)
            per_day_max[day] = val if prior is None else max(prior, val)
        rec.gear_scores = per_day_max
        new_wb.records.append(rec)
        n_rows += 1
        if rec.correct_nickname:
            n_with_truth += 1

    print(f"[2] migrated {n_rows} rows ({n_with_truth} with correct_nickname filled)")

    if dry_run:
        print("[!] --dry-run: nothing written")
        return 0

    backup = new_wb.save(backup=True)
    if backup:
        print(f"[3] legacy file backed up to {backup}")
    print(f"[4] wrote v2 workbook to {path}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--path",
        default=None,
        help="Workbook path. Defaults to data/guild_scores.xlsx.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Read & verify but do not write.",
    )
    args = parser.parse_args()
    path = Path(args.path) if args.path else (user_data_dir() / "guild_scores.xlsx")
    return migrate(path, dry_run=args.dry_run)


if __name__ == "__main__":
    raise SystemExit(main())
