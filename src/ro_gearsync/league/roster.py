"""League roster: load from — and write OCR snapshots back to —
``league_scores.xlsx``.

Per the 2026-07-02 decision the league feature keeps its OWN member list,
fully decoupled from the gear workbook. The file is user-maintained with
columns: ID / 遊戲ID / 職業 / Last_OCR_ID（OCR信心 欄已於 2026-07-07 廢除，
存在也不會再寫）. Every scan:

  1. loads this roster fresh (manual edits flow straight through),
  2. matches captured players against it (遊戲ID + Last_OCR_ID keys),
  3. writes the per-battle snapshot workbook, and
  4. updates **only** the roster's Last_OCR_ID cells for players that
     matched exactly — same self-healing loop as the gear workbook
     (a stable OCR misread becomes next week's exact-match key). Fuzzy
     matches are NOT written back until a human confirms them.

The write-back edits cells in place via openpyxl so the user's own
formatting / extra sheets survive, and a timestamped backup is taken first.
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import load_workbook

# Multi-value Last_OCR_ID helpers moved to matching/ocr_variants.py on
# 2026-07-10 (the gear workbook adopted the same format). Re-exported here
# so existing importers keep working.
from ..matching import MAX_OCR_VARIANTS, OCR_SEP, merge_ocr_variant, split_ocr_ids  # noqa: F401
from ..storage.excel import GuildScoresWorkbook, PlayerRecord
from ..utils.logging import logger

if TYPE_CHECKING:
    from .merge import BattleResult


def load_roster(roster_path: Path | None = None) -> list[PlayerRecord]:
    """Return the league roster (one :class:`PlayerRecord` per member).

    ``correct_nickname`` (遊戲ID) and ``latest_ocr_nickname`` (Last_OCR_ID)
    are the matching keys; ``player_id`` (ID) and ``profession`` (職業) ride
    along for display. The header names are shared with the gear workbook,
    so its loader works unchanged.

    A missing/empty file yields an empty roster — the caller decides
    whether to warn or abort.

    Rows whose 遊戲ID is blank are PLACEHOLDER slots (the user pre-fills
    column A with 1–150). They are kept in the returned list so the
    snapshot workbook preserves the exact row/ID layout of the roster
    file — but they never match anything (the matcher skips empty
    nicknames) and never appear as assignment candidates.
    """
    if roster_path is None:
        from ..utils.paths import league_roster_path
        roster_path = league_roster_path()
    wb = GuildScoresWorkbook.load(Path(roster_path))
    return wb.records


def roster_issues(roster: list[PlayerRecord]) -> list[str]:
    """Sanity-check the user-maintained roster; returns user-facing
    problem descriptions (empty list = OK).

    Two data problems make the Last_OCR_ID write-back silently misbehave,
    because :func:`update_roster_ocr` keys its updates on the ID column:

      * a member row (遊戲ID filled) whose ID is blank / non-numeric —
        that row can never receive a write-back;
      * duplicate IDs — several rows would all receive the same variants.

    Callers surface these at roster LOAD time so the user fixes the file
    before scanning, instead of discovering a stale Last_OCR_ID weeks later.
    """
    issues: list[str] = []
    seen: dict[int, str] = {}
    for rec in roster:
        name = rec.correct_nickname.strip()
        if not name:
            continue  # placeholder slot (blank 遊戲ID) — fine by design
        if rec.player_id is None:
            issues.append(
                f"「{name}」的 ID 欄空白或不是數字（掃描後無法回寫 Last_OCR_ID）"
            )
        elif rec.player_id in seen:
            issues.append(
                f"ID {rec.player_id} 重複：「{seen[rec.player_id]}」與「{name}」"
            )
        else:
            seen[rec.player_id] = name
    return issues


def update_roster_ocr(
    result: "BattleResult",
    roster_path: Path | None = None,
    *,
    backup: bool = True,
) -> int:
    """Write Last_OCR_ID back into the roster file, in place.

    Only rows that (a) participated this battle, (b) matched a roster
    member, and (c) did NOT need review are updated — exactly the rows we
    trust. Everything else in the workbook is left untouched (the OCR信心
    column was dropped per the 2026-07-07 spec — it saw no real use).
    Returns the number of updated members.
    """
    if roster_path is None:
        from ..utils.paths import league_roster_path
        roster_path = league_roster_path()
    roster_path = Path(roster_path)
    if not roster_path.is_file():
        logger.warning("league roster missing, skip write-back: {}", roster_path)
        return 0

    # Per member: every OCR spelling observed this battle, primary first.
    # Both battlefields contribute, and manually-assigned variants ride
    # along via ``alt_names`` — storing them is what makes next battle's
    # identical misread exact-match instead of needing assignment again.
    updates: dict[int, list[str]] = {}
    for p in result.players:
        if p.record_index is None or not p.participation or p.needs_review:
            continue
        if p.player_id is None:
            continue
        names: list[str] = []
        for stats in (p.main, p.sub):
            if stats is None:
                continue
            for name in (stats.name, *sorted(stats.alt_names)):
                if name and name not in names:
                    names.append(name)
        if names:
            updates[p.player_id] = names
    if not updates:
        return 0

    if backup:
        backup_dir = roster_path.parent / "backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(
            roster_path,
            backup_dir / f"{roster_path.stem}_{stamp}{roster_path.suffix}",
        )

    book = load_workbook(roster_path)
    ws = book[GuildScoresWorkbook.SHEET_NAME] if (
        GuildScoresWorkbook.SHEET_NAME in book.sheetnames
    ) else book[book.sheetnames[0]]
    headers = {str(c.value): i + 1 for i, c in enumerate(ws[1]) if c.value}
    id_col = headers.get("ID")
    ocr_col = headers.get("Last_OCR_ID")
    if id_col is None or ocr_col is None:
        logger.warning(
            "league roster lacks ID/Last_OCR_ID headers, skip write-back",
        )
        return 0

    n = 0
    for row in ws.iter_rows(min_row=2):
        raw_id = row[id_col - 1].value
        try:
            pid = int(raw_id) if raw_id not in (None, "") else None
        except (TypeError, ValueError):
            continue
        if pid is None or pid not in updates:
            continue
        cell_value = ws.cell(row=row[0].row, column=ocr_col).value
        # Prepend variants in reverse so the primary spelling ends up first.
        for name in reversed(updates[pid]):
            cell_value = merge_ocr_variant(cell_value, name)
        ws.cell(row=row[0].row, column=ocr_col, value=cell_value)
        n += 1

    book.save(roster_path)
    logger.info("league roster write-back: {} members updated", n)
    return n
