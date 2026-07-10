"""Compare the sheet roster against BOTH local workbooks → a reviewable
:class:`SyncPlan`.

Change kinds (one card per 編號 in the GUI/CLI; the user's decision then
drives per-file actions in :mod:`.apply`):

  join    — sheet has a name, both workbooks blank → fill 遊戲ID/職業.
  leave   — sheet slot blank/absent, workbook(s) still hold a name →
            clear 遊戲ID/職業/Last_OCR_ID (+ every gear-score column in
            guild_scores, per the 2026-07-08 decision).
  changed — both sides have a name and they differ. Ambiguous by nature:
            the user must say whether the SLOT changed hands (replace →
            clear history like leave, then fill) or the PERSON renamed
            (rename → keep gear history, clear Last_OCR_ID).
  prof    — same name, different 職業 → update 職業 only.

Comparison is exact on stripped strings — any real difference should
surface as a card for human eyes, that's the whole point of the review.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from ..storage.excel import (
    DATE_COL_RE,
    GuildScoresWorkbook,
    ID_HEADER,
    NICK_HEADER,
    OCR_CONF_HEADER,
    OCR_NICK_HEADER,
    PEAK_HEADER,
    PROFESSION_HEADER,
)
from .sheet_parse import ParsedSheet, SheetMember

KIND_JOIN = "join"
KIND_LEAVE = "leave"
KIND_CHANGED = "changed"
KIND_PROF = "prof"

KIND_LABEL = {
    KIND_JOIN: "新進成員",
    KIND_LEAVE: "退會",
    KIND_CHANGED: "遊戲ID 不同（換人或改名？）",
    KIND_PROF: "職業變更",
}

# Decisions the review UI hands to apply(): join/leave/prof cards use
# apply-or-skip; changed cards pick replace / rename / skip.
DECISION_APPLY = "apply"
DECISION_REPLACE = "replace"
DECISION_RENAME = "rename"
DECISION_SKIP = "skip"


class WorkbookLoadError(RuntimeError):
    """Display-ready failure loading one of the local workbooks."""


@dataclass
class RowState:
    row_idx: int          # 1-based worksheet row
    nickname: str
    profession: str
    peak: int | None = None       # guild only: current 最高裝評 cell value


@dataclass
class WorkbookInfo:
    path: Path
    label: str                    # short name for messages ("裝評"/"聯賽名冊")
    kind: str                     # "guild" | "league"
    sheet_name: str
    id_col: int
    nick_col: int
    prof_col: int
    ocr_col: int | None
    peak_col: int | None = None   # guild only: 最高裝評 column
    # guild only: columns wiped on leave/replace (OCR信心/最高裝評/每日裝評).
    score_cols: list[int] = field(default_factory=list)
    rows: dict[int, RowState] = field(default_factory=dict)
    excluded_ids: set[int] = field(default_factory=set)  # duplicate IDs
    warnings: list[str] = field(default_factory=list)


@dataclass
class Change:
    member_id: int
    kind: str
    sheet_name: str
    sheet_prof: str
    guild: RowState | None        # None = workbook has no row with this ID
    league: RowState | None
    # Sheet-reported 裝備評分 — seeds/raises 最高裝評 when the change is
    # applied (join fills it, replace resets it, rename raises it).
    sheet_gear: int | None = None
    notes: list[str] = field(default_factory=list)

    def summary(self) -> str:
        cur = self.guild or self.league
        cur_name = cur.nickname if cur else ""
        cur_prof = cur.profession if cur else ""
        if self.kind == KIND_JOIN:
            return f"（空位）→ {self.sheet_name}／{self.sheet_prof or '?'}"
        if self.kind == KIND_LEAVE:
            return f"{cur_name}／{cur_prof or '?'} → （清空）"
        if self.kind == KIND_CHANGED:
            return (
                f"{cur_name}／{cur_prof or '?'} → "
                f"{self.sheet_name}／{self.sheet_prof or '?'}"
            )
        return f"{cur_name}：職業 {cur_prof or '?'} → {self.sheet_prof or '?'}"


@dataclass
class PeakUpdate:
    """Sheet 裝備評分 beats the workbook's 最高裝評 for a member whose
    identity is NOT in question (names match exactly) — raise the peak."""
    member_id: int
    nickname: str
    row_idx: int                  # guild workbook row
    old: int | None
    new: int

    def summary(self) -> str:
        old = f"{self.old:,}" if self.old is not None else "（空白）"
        return f"{self.nickname}：{old} → {self.new:,}"


@dataclass
class SyncPlan:
    changes: list[Change]
    warnings: list[str]
    guild: WorkbookInfo
    league: WorkbookInfo
    sheet_member_count: int       # non-vacant slots on the sheet
    peak_updates: list[PeakUpdate] = field(default_factory=list)
    # Members whose LOCAL 最高裝評 beats the sheet's 裝備評分 (blank sheet
    # cell counts too). The tool never writes to the sheet — this only
    # feeds a "remember to update the cloud roster" reminder.
    sheet_stale_peaks: int = 0


def load_workbook_info(path: Path, label: str, kind: str) -> WorkbookInfo:
    """Read the columns roster sync cares about. ``kind == "guild"`` marks
    the gear workbook, whose score columns get wiped on leave/replace."""
    with_scores = kind == "guild"
    if not path.is_file():
        raise WorkbookLoadError(f"{label}檔案不存在：{path}")
    try:
        book = load_workbook(path, read_only=True)
    except PermissionError as exc:
        raise WorkbookLoadError(
            f"無法開啟 {path.name} — 檔案可能正被 Excel 開啟中，請先關閉。"
        ) from exc
    try:
        sheet_name = (
            GuildScoresWorkbook.SHEET_NAME
            if GuildScoresWorkbook.SHEET_NAME in book.sheetnames
            else book.sheetnames[0]
        )
        ws = book[sheet_name]
        headers: dict[str, int] = {}
        score_cols: list[int] = []
        first_row = next(ws.iter_rows(min_row=1, max_row=1), ())
        for i, cell in enumerate(first_row, start=1):
            h = str(cell.value).strip() if cell.value is not None else ""
            if not h:
                continue
            headers.setdefault(h, i)
            if with_scores and (
                h in (OCR_CONF_HEADER, PEAK_HEADER) or DATE_COL_RE.match(h)
            ):
                score_cols.append(i)

        missing = [h for h in (ID_HEADER, NICK_HEADER, PROFESSION_HEADER)
                   if h not in headers]
        if missing:
            raise WorkbookLoadError(
                f"{label}（{path.name}）缺少必要欄位：{'、'.join(missing)}"
            )

        info = WorkbookInfo(
            path=path,
            label=label,
            kind=kind,
            sheet_name=sheet_name,
            id_col=headers[ID_HEADER],
            nick_col=headers[NICK_HEADER],
            prof_col=headers[PROFESSION_HEADER],
            ocr_col=headers.get(OCR_NICK_HEADER),
            peak_col=headers.get(PEAK_HEADER) if with_scores else None,
            score_cols=score_cols,
        )

        # enumerate for the row number — read-only mode hands out EmptyCell
        # objects that carry no .row attribute.
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            raw_id = row[info.id_col - 1].value if len(row) >= info.id_col else None
            try:
                mid = int(raw_id) if raw_id not in (None, "") else None
            except (TypeError, ValueError):
                continue
            if mid is None:
                continue

            def _text(col: int) -> str:
                v = row[col - 1].value if len(row) >= col else None
                return str(v).strip() if v is not None else ""

            peak = None
            if info.peak_col is not None:
                raw_peak = _text(info.peak_col).replace(",", "")
                try:
                    peak = int(float(raw_peak)) if raw_peak else None
                except ValueError:
                    peak = None

            state = RowState(
                row_idx=row_idx,
                nickname=_text(info.nick_col),
                profession=_text(info.prof_col),
                peak=peak,
            )
            if mid in info.rows:
                if mid not in info.excluded_ids:
                    info.excluded_ids.add(mid)
                    info.warnings.append(
                        f"{label}（{path.name}）ID {mid} 重複 — 此 ID 不同步，"
                        "請先修正檔案。"
                    )
                continue
            info.rows[mid] = state
        for mid in info.excluded_ids:
            info.rows.pop(mid, None)
        return info
    finally:
        book.close()


def compute_plan(
    parsed: ParsedSheet,
    guild_path: Path | None = None,
    league_path: Path | None = None,
) -> SyncPlan:
    """Diff the parsed sheet against both workbooks."""
    from ..utils.paths import default_workbook_path, league_roster_path

    guild = load_workbook_info(
        guild_path or default_workbook_path(), "裝評工作簿", "guild",
    )
    league = load_workbook_info(
        league_path or league_roster_path(), "聯賽名冊", "league",
    )

    warnings = list(parsed.warnings) + guild.warnings + league.warnings
    skip_ids = parsed.excluded_ids | guild.excluded_ids | league.excluded_ids

    all_ids = sorted(
        (set(parsed.members) | set(guild.rows) | set(league.rows)) - skip_ids
    )
    changes: list[Change] = []
    peak_updates: list[PeakUpdate] = []
    sheet_stale_peaks = 0
    for mid in all_ids:
        sheet: SheetMember | None = parsed.members.get(mid)
        s_name = sheet.nickname if sheet else ""
        s_prof = sheet.profession if sheet else ""
        s_gear = sheet.gear_score if sheet else None
        g = guild.rows.get(mid)
        l = league.rows.get(mid)

        # Sheet 裝評 vs 最高裝評 — high-water mark, raise only. Restricted
        # to rows whose identity is beyond doubt (names match exactly);
        # join/replace/rename rows get their peak via the change itself.
        if (
            s_gear
            and s_name
            and g is not None
            and guild.peak_col is not None
            and g.nickname == s_name
            and s_gear > (g.peak or 0)
        ):
            peak_updates.append(
                PeakUpdate(mid, s_name, g.row_idx, g.peak, s_gear)
            )

        # Opposite direction — local peak beats the sheet (a blank sheet
        # cell counts: the cloud is missing the number entirely). Counted
        # only, for the "update the cloud roster" reminder.
        if (
            parsed.has_gear_column
            and s_name
            and g is not None
            and g.nickname == s_name
            and (g.peak or 0) > (s_gear or 0)
        ):
            sheet_stale_peaks += 1

        notes: list[str] = []
        if s_name:
            if g is None:
                notes.append(f"裝評工作簿沒有 ID {mid} 的列 — 該檔跳過，請自行補列")
            if l is None:
                notes.append(f"聯賽名冊沒有 ID {mid} 的列 — 該檔跳過，請自行補列")
            if g is None and l is None:
                warnings.append(
                    f"試算表編號 {mid}（{s_name}）在兩份 Excel 都沒有對應的 ID 列，"
                    "無法同步 — 請先在 Excel 補上該列。"
                )
                continue

        names = [st.nickname for st in (g, l) if st is not None]
        if g and l and (g.nickname != l.nickname or g.profession != l.profession):
            notes.append(
                f"兩份 Excel 目前不一致（裝評：{g.nickname or '空'}／"
                f"{g.profession or '?'}，聯賽：{l.nickname or '空'}／"
                f"{l.profession or '?'}）"
            )

        has_excel_name = any(names)
        if s_name:
            if any(n and n != s_name for n in names):
                kind = KIND_CHANGED
            elif not has_excel_name:
                kind = KIND_JOIN
            elif any(st is not None and not st.nickname for st in (g, l)):
                # One workbook already has the member, the other is blank —
                # a fill-in; join semantics fit (nothing destructive).
                kind = KIND_JOIN
                notes.append("其中一份 Excel 已有此成員，僅補齊另一份")
            elif s_prof and any(
                st is not None and st.nickname and st.profession != s_prof
                for st in (g, l)
            ):
                kind = KIND_PROF
            else:
                continue  # in sync
        else:
            if not has_excel_name:
                continue  # vacant on both sides
            kind = KIND_LEAVE

        changes.append(Change(
            member_id=mid, kind=kind,
            sheet_name=s_name, sheet_prof=s_prof,
            guild=g, league=l, sheet_gear=s_gear, notes=notes,
        ))

    return SyncPlan(
        changes=changes,
        warnings=warnings,
        guild=guild,
        league=league,
        sheet_member_count=sum(1 for m in parsed.members.values() if m.nickname),
        peak_updates=peak_updates,
        sheet_stale_peaks=sheet_stale_peaks,
    )
