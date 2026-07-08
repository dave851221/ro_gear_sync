"""Write one league battle to a fresh snapshot workbook.

Per the 2026-07-02 refinements: a **single wide sheet** (「聯賽戰績」),
one row per roster member, with

  ID | 遊戲ID | 職業 | 參與 | 主-…(8 欄) | 副-…(12 欄)

* 參與 records 主 / 副 / 主+副 (blank = didn't play).
* Every screen's metrics get their own column — 主戰場 contributes the
  輸出+輔助 eight, 副戰場 adds the 戰略 four on top.
* Unmatched captures append at the bottom as red rows (OCR name in the
  遊戲ID column) for manual handling.

Every scan produces a **new** standalone file named
``league_scores_YYYYMMDD_HHMMSS.xlsx`` — no accumulation, and the gear
workbook is never touched. Styling mirrors the gear workbook (storage/excel.py).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .merge import BattleResult
from .model import (
    BATTLEFIELD_LABEL,
    DPS_METRICS,
    LABELS,
    STRATEGY_METRICS,
    SUPPORT_METRICS,
)

SHEET_NAME = "聯賽戰績"

# Per-battlefield metric layouts (main has no 戰略 view).
MAIN_METRICS: tuple[str, ...] = DPS_METRICS + SUPPORT_METRICS
SUB_METRICS: tuple[str, ...] = DPS_METRICS + SUPPORT_METRICS + STRATEGY_METRICS

META_HEADERS: tuple[str, ...] = ("ID", "遊戲ID", "職業", "參與")
HEADERS: tuple[str, ...] = (
    META_HEADERS
    + tuple(f"主-{LABELS[k]}" for k in MAIN_METRICS)
    + tuple(f"副-{LABELS[k]}" for k in SUB_METRICS)
)
_PARTICIPATION_COL = META_HEADERS.index("參與") + 1

# Fills.
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_USER_FILL = PatternFill("solid", fgColor="FFF2CC")    # 遊戲ID / 職業
_MAIN_FILL = PatternFill("solid", fgColor="C6EFCE")    # 主 (green)
_SUB_FILL = PatternFill("solid", fgColor="BDD7EE")     # 副 (blue)
_BOTH_FILL = PatternFill("solid", fgColor="92D050")    # 主+副 (dark green)
_RED_FILL = PatternFill("solid", fgColor="FFC7CE")     # review / unmatched
_MAIN_NUM_FILL = PatternFill("solid", fgColor="E2EFDA")  # 主-metric cells
_SUB_NUM_FILL = PatternFill("solid", fgColor="DDEBF7")   # 副-metric cells

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(border_style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def league_filename(when: datetime | None = None) -> str:
    # Seconds included: the review dialog encourages "adjust → write
    # again", and two presses within the same minute must yield two
    # files (minute granularity silently overwrote the first, and hit
    # PermissionError if the user had it open in Excel).
    when = when or datetime.now()
    return f"league_scores_{when:%Y%m%d_%H%M%S}.xlsx"


def _participation_fill(part: str) -> PatternFill | None:
    return {
        BATTLEFIELD_LABEL["main"]: _MAIN_FILL,
        BATTLEFIELD_LABEL["sub"]: _SUB_FILL,
        f"{BATTLEFIELD_LABEL['main']}+{BATTLEFIELD_LABEL['sub']}": _BOTH_FILL,
    }.get(part)


def _write_sheet(ws: Worksheet, result: BattleResult) -> None:
    ws.append(list(HEADERS))
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER

    n_meta = len(META_HEADERS)
    n_main = len(MAIN_METRICS)

    for r_idx, p in enumerate(result.players, start=2):
        main = p.main.as_metric_dict() if p.main is not None else {}
        sub = p.sub.as_metric_dict() if p.sub is not None else {}
        values: list = [p.player_id, p.nickname, p.profession, p.participation]
        values += [main.get(k) if p.main is not None else None for k in MAIN_METRICS]
        values += [sub.get(k) if p.sub is not None else None for k in SUB_METRICS]

        is_red = p.is_unmatched or p.needs_review
        for c_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = _BORDER
            cell.alignment = _CENTER
            is_metric = c_idx > n_meta
            if is_metric and isinstance(value, int):
                cell.number_format = "#,##0"
            if is_red:
                cell.fill = _RED_FILL
                continue
            header = HEADERS[c_idx - 1]
            if header in ("遊戲ID", "職業"):
                cell.fill = _USER_FILL
            elif header == "參與":
                fill = _participation_fill(p.participation)
                if fill:
                    cell.fill = fill
            elif is_metric and value is not None:
                cell.fill = (
                    _MAIN_NUM_FILL if c_idx <= n_meta + n_main else _SUB_NUM_FILL
                )
        # Review note in the spill column on the far right.
        if p.review_note:
            note = ws.cell(row=r_idx, column=len(HEADERS) + 1, value=p.review_note)
            note.alignment = Alignment(horizontal="left", vertical="center")
            note.font = Font(color="9C0006")

    widths = {"ID": 6, "遊戲ID": 22, "職業": 12, "參與": 9}
    for i, h in enumerate(HEADERS, start=1):
        letter = get_column_letter(i)
        if h in widths:
            ws.column_dimensions[letter].width = widths[h]
        else:
            # Metric columns: wide enough for "主-王的最後一擊" style headers
            # and 9-digit numbers with separators.
            ws.column_dimensions[letter].width = 13
    ws.column_dimensions[get_column_letter(len(HEADERS) + 1)].width = 32

    # Freeze the meta block; filter over the full table.
    ws.freeze_panes = f"{get_column_letter(n_meta + 1)}2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(HEADERS))}{len(result.players) + 1}"
    )


def write_battle(
    result: BattleResult,
    output_dir: Path | None = None,
    *,
    when: datetime | None = None,
) -> Path:
    """Write ``result`` to a fresh ``league_scores_*.xlsx`` and return its path."""
    if output_dir is None:
        from ..utils.paths import league_output_dir
        output_dir = league_output_dir()
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / league_filename(when)

    book = Workbook()
    ws = book.active
    ws.title = SHEET_NAME
    _write_sheet(ws, result)
    book.save(out_path)
    return out_path
