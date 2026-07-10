"""Apply user-confirmed roster changes to both workbooks, IN PLACE.

Writes go through openpyxl cell edits (never GuildScoresWorkbook.save(),
whose full rebuild would wipe the user's own formatting/extra sheets) —
same pattern as the league Last_OCR_ID write-back. Both files are backed
up to ``<file dir>/backups/`` before saving.

Decision → per-workbook actions (only rows whose current state actually
needs the edit are touched; a workbook already in sync is left alone):

  join    → fill 遊戲ID/職業 into the blank row.
  leave   → clear 遊戲ID/職業/Last_OCR_ID; guild also clears every
            score column (OCR信心/最高裝評/per-day columns).
  replace — the slot changed hands: clear like leave (incl. guild score
            history — those numbers belong to the previous person),
            then fill the new 遊戲ID/職業.
  rename  — same person, new name: set 遊戲ID/職業, clear Last_OCR_ID
            (old OCR variants can't match the new name), KEEP scores.
  prof    → update 職業 only.

The sheet's 裝備評分 column feeds 最高裝評 (guild workbook only, 2026-07-10):

  * join/replace seed the peak with the sheet value (fresh identity);
    rename raises it only if the sheet value is bigger;
  * ``sync_peaks=True`` additionally applies ``plan.peak_updates`` — the
    raise-only updates for members whose names already match. The peak
    cell survives GuildScoresWorkbook round-trips because load() keeps
    the stored peak whenever it beats every per-day column.
"""
from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..utils.logging import logger
from .diff import (
    Change,
    DECISION_APPLY,
    DECISION_RENAME,
    DECISION_REPLACE,
    DECISION_SKIP,
    KIND_CHANGED,
    KIND_JOIN,
    KIND_LEAVE,
    KIND_PROF,
    SyncPlan,
    WorkbookInfo,
)


class ApplyError(RuntimeError):
    """Display-ready apply failure (typically the file is open in Excel)."""


@dataclass
class ApplyReport:
    applied: list[str] = field(default_factory=list)   # one line per change
    skipped: int = 0
    backups: list[Path] = field(default_factory=list)
    peak_applied: list[str] = field(default_factory=list)  # one line per peak

    @property
    def applied_count(self) -> int:
        return len(self.applied)


def _backup(path: Path) -> Path:
    backup_dir = path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = backup_dir / f"{path.stem}_{stamp}{path.suffix}"
    shutil.copy2(path, dest)
    return dest


def _apply_to_workbook(
    ws: Worksheet, info: WorkbookInfo, change: Change, decision: str,
) -> bool:
    """Edit one workbook's row for this change. Returns True if any cell
    was written."""
    state = change.guild if info.kind == "guild" else change.league
    if state is None:
        return False
    row = state.row_idx

    def _set(col: int | None, value) -> None:
        # Assign via .value — ws.cell(..., value=None) treats None as
        # "no value given" and silently skips the write, so clears would
        # be no-ops through the constructor form.
        if col is not None:
            ws.cell(row=row, column=col).value = value

    def _clear_identity() -> None:
        _set(info.nick_col, None)
        _set(info.prof_col, None)
        _set(info.ocr_col, None)

    def _clear_scores() -> None:
        for col in info.score_cols:
            _set(col, None)

    def _fill(name: str, prof: str) -> None:
        _set(info.nick_col, name or None)
        _set(info.prof_col, prof or None)

    def _seed_peak(*, raise_only: bool) -> None:
        """Write the sheet's 裝備評分 into 最高裝評 alongside an identity
        fill. join/replace = fresh identity, take the sheet value as-is;
        rename (raise_only) = same person, high-water mark semantics."""
        if info.kind != "guild" or info.peak_col is None or not change.sheet_gear:
            return
        if raise_only and change.sheet_gear <= (state.peak or 0):
            return
        _set(info.peak_col, change.sheet_gear)

    if change.kind == KIND_LEAVE and decision == DECISION_APPLY:
        if not state.nickname:
            return False
        _clear_identity()
        _clear_scores()
        return True

    if change.kind == KIND_JOIN and decision == DECISION_APPLY:
        if state.nickname == change.sheet_name:
            # Already present here — at most a profession touch-up.
            if change.sheet_prof and state.profession != change.sheet_prof:
                _set(info.prof_col, change.sheet_prof)
                return True
            return False
        if state.nickname:
            # A different name in a "join" slot means the plan went stale
            # since it was computed — leave the row alone.
            return False
        _fill(change.sheet_name, change.sheet_prof)
        _seed_peak(raise_only=False)
        return True

    if change.kind == KIND_CHANGED and decision in (DECISION_REPLACE, DECISION_RENAME):
        if state.nickname == change.sheet_name:
            if change.sheet_prof and state.profession != change.sheet_prof:
                _set(info.prof_col, change.sheet_prof)
                return True
            return False
        _set(info.ocr_col, None)
        if decision == DECISION_REPLACE:
            _clear_scores()
        _fill(change.sheet_name, change.sheet_prof)
        _seed_peak(raise_only=decision == DECISION_RENAME)
        return True

    if change.kind == KIND_PROF and decision == DECISION_APPLY:
        if state.nickname and state.profession != change.sheet_prof:
            _set(info.prof_col, change.sheet_prof or None)
            return True
        return False

    return False


_DECISION_LABEL = {
    DECISION_APPLY: "套用",
    DECISION_REPLACE: "換人（清空舊資料後填入）",
    DECISION_RENAME: "同一人改名/換職業（保留裝評紀錄）",
}


def apply_plan(
    plan: SyncPlan,
    decisions: dict[int, str],
    *,
    sync_peaks: bool = False,
) -> ApplyReport:
    """``decisions`` maps 編號 → apply/replace/rename/skip. Anything not
    in the dict counts as skip. ``sync_peaks`` additionally writes
    ``plan.peak_updates`` into the guild workbook's 最高裝評 column.
    Both workbooks are opened first (fail early if either is locked),
    backed up, edited, then saved."""
    report = ApplyReport()
    todo = [
        (c, decisions.get(c.member_id, DECISION_SKIP))
        for c in plan.changes
    ]
    report.skipped = sum(1 for _, d in todo if d == DECISION_SKIP)
    todo = [(c, d) for c, d in todo if d != DECISION_SKIP]
    peaks = plan.peak_updates if sync_peaks else []
    if not todo and not peaks:
        return report

    books = []
    try:
        for info in (plan.guild, plan.league):
            try:
                book = load_workbook(info.path)
            except PermissionError as exc:
                raise ApplyError(
                    f"無法開啟 {info.path.name} — 檔案可能正被 Excel 開啟中，"
                    "請先關閉後再套用。"
                ) from exc
            ws = (
                book[info.sheet_name]
                if info.sheet_name in book.sheetnames
                else book[book.sheetnames[0]]
            )
            books.append((info, book, ws))

        for change, decision in todo:
            touched = [
                info.label
                for info, _book, ws in books
                if _apply_to_workbook(ws, info, change, decision)
            ]
            where = "、".join(touched) if touched else "（兩份檔案皆已是最新，未變動）"
            report.applied.append(
                f"ID {change.member_id}｜{change.summary()}｜"
                f"{_DECISION_LABEL.get(decision, decision)} → {where}"
            )

        if peaks:
            g_info, _g_book, g_ws = next(
                (info, book, ws) for info, book, ws in books
                if info.kind == "guild"
            )
            for pu in peaks:
                # Staleness guard — same spirit as the join branch: if the
                # row's name no longer matches what the plan saw, skip it.
                cur = g_ws.cell(row=pu.row_idx, column=g_info.nick_col).value
                cur_name = str(cur).strip() if cur is not None else ""
                if cur_name != pu.nickname:
                    report.peak_applied.append(
                        f"ID {pu.member_id}｜{pu.summary()}｜"
                        "略過（該列遊戲ID 已變動）"
                    )
                    continue
                g_ws.cell(row=pu.row_idx, column=g_info.peak_col).value = pu.new
                report.peak_applied.append(f"ID {pu.member_id}｜{pu.summary()}")

        # All edits staged in memory — back up, then save both.
        for info, _book, _ws in books:
            report.backups.append(_backup(info.path))
        for info, book, _ws in books:
            try:
                book.save(info.path)
            except PermissionError as exc:
                raise ApplyError(
                    f"寫入 {info.path.name} 失敗 — 檔案可能正被 Excel 開啟中。"
                    "已寫入的備份在 backups/ 資料夾。"
                ) from exc
        logger.info(
            "roster sync applied: {} changes, {} skipped, {} peak updates",
            len(todo), report.skipped, len(report.peak_applied),
        )
        return report
    finally:
        for _info, book, _ws in books:
            book.close()
